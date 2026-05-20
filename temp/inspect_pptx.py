from pptx import Presentation

prs = Presentation('/Users/vatsravi/2026/AnyNewProject/sample-travel-concierge-with-amazon-bedrock-agentcore-and-nova-sonic/telephony-ai-host.pptx')
slide = prs.slides[0]
print('Slide size:', prs.slide_width/914400, 'x', prs.slide_height/914400, 'inches')
print('Number of shapes:', len(slide.shapes))
for i, s in enumerate(slide.shapes):
    name = s.name
    left = round(s.left/914400, 2)
    top = round(s.top/914400, 2)
    w = round(s.width/914400, 2)
    h = round(s.height/914400, 2)
    stype = s.shape_type
    text = ''
    if hasattr(s, 'text') and s.text:
        text = s.text[:60].replace('\n', ' ')
    print(f'  {i}: type={stype} name={name!r} pos=({left},{top}) size=({w},{h}) text={text!r}')
